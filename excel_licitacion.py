"""Excel de trabajo de una licitación: pestañas Info, Solvencia, C.Adjudicación y Sobres.

Lo rellena preparar_carpeta.py con (1) los datos del monitor (expediente, órgano, importe, enlace, fecha
límite), (2) el bloque «datos_excel» que la IA añade a la Revisión (ver PROMPT_SISTEMA_INFORME) y (3) los
sobres y documentos del Bloque 5. Las casillas son las casillas nativas de Excel 365 (valor VERDADERO/FALSO);
en versiones antiguas de Excel o en LibreOffice se ven como VERDADERO/FALSO, que se puede teclear igual.

Lo que no consta en el pliego se deja en amarillo con «REVISAR».
"""
import math
import re
import shutil
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FUENTE = "Aptos Narrow"        # la de la plantilla de CAI
AZUL = "9FD8EE"                # cabeceras de bloque (como la plantilla)
AZUL_SUAVE = "CFECF7"          # cabeceras de criterio
TEAL = "07AAAA"                # color de CAI Consultores
NARANJA = "FBE3D6"             # Posible UTE
AMARILLO = "FFF2A8"            # dato que no consta: revisar
GRIS = "EDEDED"
GRIS_CLARO = "F6F6F6"
TEXTO_REVISAR = "REVISAR: no consta en el pliego"
FORMATO_EUR = '#,##0.00\\ "€"'

_L = Side(style="thin", color="A6A6A6")
BORDE = Border(left=_L, right=_L, top=_L, bottom=_L)

TIPOS_CONTRATO = {"1": "Suministros", "2": "Servicios", "3": "Obras", "21": "Gestión de Servicios Públicos",
                  "22": "Concesión de Servicios", "31": "Concesión de Obras Públicas", "32": "Concesión de Obras",
                  "40": "Colaboración público-privada", "7": "Administrativo especial", "8": "Privado", "50": "Patrimonial"}
PROCEDIMIENTOS = {"1": "Abierto", "9": "Abierto simplificado", "2": "Restringido", "3": "Negociado sin publicidad",
                  "4": "Negociado con publicidad", "5": "Diálogo competitivo", "6": "Contrato menor",
                  "7": "Derivado de acuerdo marco", "8": "Concurso de proyectos", "10": "Asociación para la innovación",
                  "13": "Licitación con negociación", "100": "Normas internas", "999": "Otros"}
ETIQUETA_CAI = {"cumple": "Cumple", "parcial": "Parcial", "no": "No cumple", "no_consta": "Sin datos"}


# ------------------------------------------------------------------ utilidades
def _txt(x):
    if x is None:
        return None
    t = str(x).strip()
    return t or None


def _num(x):
    if x is None or isinstance(x, bool) or x == "":
        return None
    try:
        n = float(str(x).replace(",", ".")) if isinstance(x, str) else float(x)
    except ValueError:
        return None
    return int(n) if n == int(n) else n


def _lista(x):
    return [e for e in x if isinstance(e, dict)] if isinstance(x, list) else []


def _alto(*pares, minimo=18.0):
    """Altura de fila estimada para texto con ajuste de línea: (texto, ancho de columna) por celda."""
    lineas = 1
    for texto, ancho in pares:
        if not texto:
            continue
        n = sum(max(1, math.ceil(len(p) / max(ancho * 1.15, 1))) for p in str(texto).split("\n"))
        lineas = max(lineas, n)
    return min(409.0, max(minimo, 15.0 * lineas + 4))


def _fmt(c, *, negrita=False, relleno=None, color=None, h="left", v="center", wrap=True, tam=11,
         cursiva=False, borde=True, num=None, shrink=False, subrayado=None):
    c.font = Font(name=FUENTE, size=tam, bold=negrita, italic=cursiva, color=color, underline=subrayado)
    if relleno:
        c.fill = PatternFill("solid", fgColor=relleno)
    c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=(wrap and not shrink) or None, shrink_to_fit=shrink or None)
    if borde:
        c.border = BORDE
    if num:
        c.number_format = num


def _escribe(ws, ref, valor=None, **kw):
    c = ws[ref]
    if valor is not None:
        c.value = valor
    _fmt(c, **kw)
    return c


def _combina(ws, rango, valor=None, **kw):
    for fila in ws[rango]:
        for c in fila:
            _fmt(c, **kw)
    ws[rango.split(":")[0]].value = valor
    ws.merge_cells(rango)


def _casilla(ws, ref, marcada=False, relleno=None):
    # shrink_to_fit marca la celda como «casilla»: _activar_casillas() convierte esos estilos
    return _escribe(ws, ref, bool(marcada), h="center", shrink=True, relleno=relleno)


def _revisar(ws, rango, **kw):
    if ":" in rango:
        _combina(ws, rango, TEXTO_REVISAR, relleno=AMARILLO, cursiva=True, **kw)
    else:
        _escribe(ws, rango, TEXTO_REVISAR, relleno=AMARILLO, cursiva=True, **kw)


def _anchos(ws, anchos):
    for col, w in anchos.items():
        ws.column_dimensions[col].width = w


def _preparar_hoja(ws, orientacion="landscape"):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = orientacion
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _parse_fecha(*valores):
    for v in valores:
        s = _txt(v)
        if not s:
            continue
        try:
            return datetime.strptime(s[:10], "%Y-%m-%d")
        except ValueError:
            continue
    return None


# ------------------------------------------------------------------ Info
def _hoja_info(wb, registro, datos, exp, hoy):
    ws = wb.active
    ws.title = "Info"
    _preparar_hoja(ws, "portrait")
    _anchos(ws, {"A": 34, "B": 36, "C": 36})
    info = datos.get("info") if isinstance(datos.get("info"), dict) else {}
    _combina(ws, "A1:C1", f"Ficha de la licitación · {exp}", negrita=True, tam=14, color=TEAL, borde=False)
    _combina(ws, "A2:C2", f"Generada el {hoy:%d/%m/%Y} a partir de la Revisión IA del monitor. Contrasta siempre con los "
                          "pliegos oficiales. En amarillo: dato que no consta, a revisar.",
             cursiva=True, tam=9, borde=False, color="595959")
    ws.row_dimensions[2].height = 26
    r = 4

    def etiqueta(fila, texto, filas=1):
        rango = f"A{fila}:A{fila + filas - 1}"
        _combina(ws, rango, texto, negrita=True, relleno=AZUL, h="right") if filas > 1 else \
            _escribe(ws, f"A{fila}", texto, negrita=True, relleno=AZUL, h="right")

    def texto(etiq, valor, extra=None):
        nonlocal r
        etiqueta(r, etiq)
        v = _txt(valor)
        if v:
            _combina(ws, f"B{r}:C{r}", v)
            ws.row_dimensions[r].height = _alto((v, 70), minimo=22)
        else:
            _revisar(ws, f"B{r}:C{r}")
            ws.row_dimensions[r].height = 22
        r += 1

    texto("Nº Expediente", registro.get("folder_id"))
    texto("Nombre de la licitación", registro.get("descripcion") or registro.get("titulo"))
    texto("Administración contratante", registro.get("organo"))

    etiqueta(r, "Enlace a la plataforma")
    enlace = _txt(registro.get("link"))
    if enlace:
        _combina(ws, f"B{r}:C{r}", "Plataforma de Contratación del Sector Público (ficha de la licitación)",
                 color="0563C1", subrayado="single")
        ws[f"B{r}"].hyperlink = enlace
    else:
        _revisar(ws, f"B{r}:C{r}")
    ws.row_dimensions[r].height = 22
    r += 1

    cod_tipo = _txt(registro.get("tipo_contrato"))
    texto("Tipo de contrato", TIPOS_CONTRATO.get(cod_tipo, cod_tipo))
    cod_proc = _txt(registro.get("procedimiento"))
    texto("Procedimiento", PROCEDIMIENTOS.get(cod_proc, cod_proc))
    cpv = registro.get("cpv_codes")
    texto("CPV", ", ".join(cpv) if isinstance(cpv, list) else cpv)
    texto("División en lotes", info.get("lotes"))

    # IVA y precios: «Con IVA» se calcula con el IVA de esta celda (21 % salvo que el pliego diga otro)
    etiqueta(r, "IVA aplicado (%)")
    iva = _num(info.get("iva_pct"))
    _escribe(ws, f"B{r}", iva if iva is not None else 21, h="left", num="0.##")
    _escribe(ws, f"C{r}", "" if iva is not None else "Supuesto: el pliego no lo indica (revisar)", cursiva=True,
             tam=9, relleno=None if iva is not None else AMARILLO)
    fila_iva = r
    r += 1

    def importe(etiq, valor):
        nonlocal r
        etiqueta(r, etiq, 2)
        _escribe(ws, f"B{r}", "Sin IVA", negrita=True, relleno=AZUL_SUAVE, h="center")
        _escribe(ws, f"C{r}", "Con IVA", negrita=True, relleno=AZUL_SUAVE, h="center")
        if valor is not None:
            _escribe(ws, f"B{r + 1}", valor, num=FORMATO_EUR, h="center")
            _escribe(ws, f"C{r + 1}", f"=B{r + 1}*(1+$B${fila_iva}/100)", num=FORMATO_EUR, h="center")
        else:
            _revisar(ws, f"B{r + 1}:C{r + 1}")
        r += 2

    pbl = _num(registro.get("importe_num"))
    importe("Precio base de la licitación", pbl if pbl is not None else _num(registro.get("importe")))
    importe("Valor estimado del contrato", _num(info.get("valor_estimado")))

    texto("Plazo de ejecución", info.get("plazo_ejecucion"))

    etiqueta(r, "Fecha de presentación")
    fecha = _parse_fecha(registro.get("fecha_limite"), info.get("fecha_presentacion"))
    if fecha:
        _escribe(ws, f"B{r}", fecha, num="dd/mm/yyyy", h="center")
        _escribe(ws, f"C{r}", f'=IF(B{r}<TODAY(),"Plazo vencido","Faltan "&(B{r}-TODAY())&" días")', h="center", negrita=True)
    else:
        _revisar(ws, f"B{r}:C{r}")
    ws.row_dimensions[r].height = 22
    r += 1

    texto("Consultas", info.get("consultas"))
    texto("Garantía provisional", info.get("garantia_provisional"))
    texto("Garantía definitiva", info.get("garantia_definitiva"))
    texto("Seguro de responsabilidad civil", info.get("seguro_rc"))
    texto("Plazo de garantía", info.get("plazo_garantia"))
    texto("Revisión de precios", info.get("revision_precios"))
    texto("UTE", info.get("ute"))
    texto("Subcontratación", info.get("subcontratacion"))
    texto("Umbral mínimo de puntuación técnica", info.get("umbral_tecnico") or datos.get("umbral"))


# ------------------------------------------------------------------ Solvencia
def _hoja_solvencia(wb, datos, exp):
    ws = wb.create_sheet("Solvencia")
    _preparar_hoja(ws)
    A, B, C, E, G = 44, 55, 62, 30, 30
    _anchos(ws, {"A": A, "B": B, "C": C, "D": 7, "E": E, "F": 7, "G": G})
    sol = datos.get("solvencia") if isinstance(datos.get("solvencia"), dict) else {}
    nota = _txt(sol.get("nota")) or "Requisitos mínimos de solvencia y de equipo según el pliego."
    _combina(ws, "A1:C1", nota, h="center", borde=False, tam=10, cursiva=True)
    _combina(ws, "D1:E1", "CAI Consultores", negrita=True, relleno=TEAL, color="FFFFFF", h="center")
    _combina(ws, "F1:G1", "Posible UTE", negrita=True, relleno=NARANJA, h="center")
    ws.row_dimensions[1].height = _alto((nota, A + B + C), minimo=30)
    ws.freeze_panes = "A2"
    r = 2

    def seccion(titulo, cab_b, cab_c, items):
        nonlocal r
        _escribe(ws, f"A{r}", titulo, negrita=True, relleno=AZUL)
        _escribe(ws, f"B{r}", cab_b, negrita=True, relleno=AZUL)
        _escribe(ws, f"C{r}", cab_c, negrita=True, relleno=AZUL)
        for col in "DEFG":
            _escribe(ws, f"{col}{r}", None, relleno=AZUL)
        r += 1
        if not items:
            _revisar(ws, f"A{r}:C{r}")
            for col in "DF":
                _casilla(ws, f"{col}{r}")
            for col in "EG":
                _escribe(ws, f"{col}{r}")
            r += 1
            return
        for it in items:
            nombre, valor, acred = _txt(it.get("nombre")), _txt(it.get("valor")), _txt(it.get("acreditacion"))
            _escribe(ws, f"A{r}", nombre or "")
            _escribe(ws, f"B{r}", valor or "")
            _escribe(ws, f"C{r}", acred or "")
            cai = _txt(it.get("cai"))
            _casilla(ws, f"D{r}", cai == "cumple")
            nota_cai = _txt(it.get("cai_nota"))
            etiq = ETIQUETA_CAI.get(cai)
            _escribe(ws, f"E{r}", (f"{etiq}: {nota_cai}" if etiq and nota_cai else etiq or nota_cai or ""), tam=10)
            _casilla(ws, f"F{r}")
            _escribe(ws, f"G{r}", "", tam=10)
            ws.row_dimensions[r].height = _alto((nombre, A), (valor, B), (acred, C), (ws[f"E{r}"].value, E * 1.1), minimo=30)
            r += 1

    seccion("Solvencia económica", "Valor", "Acreditación", _lista(sol.get("economica")))
    seccion("Solvencia técnica (empresa)", "Requisito", "Acreditación", _lista(sol.get("tecnica")))
    seccion("Equipo mínimo", "Titulación y experiencia", "Acreditación", _lista(sol.get("equipo")))


# ------------------------------------------------------------------ C.Adjudicación
def _hoja_adjudicacion(wb, datos, exp):
    ws = wb.create_sheet("C.Adjudicación")
    _preparar_hoja(ws)
    anch = {"A": 24, "B": 58, "C": 11, "D": 9, "E": 9, "F": 16, "G": 11, "H": 7, "I": 32, "J": 7, "K": 32}
    _anchos(ws, anch)
    ancho_total = sum(anch.values())
    _combina(ws, "A1:B1", f"Criterios de adjudicación · {exp}", negrita=True, tam=14, color=TEAL, borde=False)
    _combina(ws, "H1:I1", "CAI Consultores", negrita=True, relleno=TEAL, color="FFFFFF", h="center")
    _combina(ws, "J1:K1", "Posible UTE", negrita=True, relleno=NARANJA, h="center")
    ws.row_dimensions[1].height = 30
    for col, texto in zip("ABCDEFGHIJK", ["Criterio", "Apartado", "Límite de puntos", "Págs. A4", "Págs. A3", "Responsable",
                                          "Ptos aprox.", "✓", "Proyectos / perfiles de CAI", "✓", "Proyectos / perfiles de la UTE"]):
        _escribe(ws, f"{col}2", texto, negrita=True, relleno=AZUL, h="center")
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = "A3"
    rojo = Font(name=FUENTE, color="C00000", bold=True)
    cabeceras_puntos = []   # filas de cabecera de criterio: suman el total del pliego
    r = 3

    def titulo_bloque(texto):
        nonlocal r
        _combina(ws, f"A{r}:K{r}", texto, negrita=True, relleno=AZUL, tam=12)
        ws.row_dimensions[r].height = 22
        r += 1
        return r - 1

    def cabecera_criterio(nombre, puntos, a4=None, a3=None):
        nonlocal r
        _combina(ws, f"A{r}:B{r}", nombre or "(sin nombre)", negrita=True, relleno=AZUL_SUAVE)
        _escribe(ws, f"C{r}", puntos, negrita=True, relleno=AZUL_SUAVE, h="center", color="C00000")
        _escribe(ws, f"D{r}", a4, relleno=AZUL_SUAVE, h="center")
        _escribe(ws, f"E{r}", a3, relleno=AZUL_SUAVE, h="center")
        for col in "FGHIJK":
            _escribe(ws, f"{col}{r}", None, relleno=AZUL_SUAVE)
        ws.row_dimensions[r].height = _alto((nombre, anch["A"] + anch["B"]), minimo=22)
        cabeceras_puntos.append(r)
        r += 1
        return r - 1

    def nota(texto):
        nonlocal r
        if not texto:
            return
        _combina(ws, f"A{r}:K{r}", texto, cursiva=True, tam=10, relleno=GRIS_CLARO)
        ws.row_dimensions[r].height = _alto((texto, ancho_total), minimo=18)
        r += 1

    def apartados(cab, filas, paginas, extra_vacias=0):
        """filas: lista de (nombre, puntos). Devuelve la fila de suma."""
        nonlocal r
        filas = list(filas) + [("", None)] * extra_vacias
        primera, ultima = r, r + len(filas) - 1
        for nombre, puntos in filas:
            _escribe(ws, f"A{r}")
            _escribe(ws, f"B{r}", nombre or "")
            _escribe(ws, f"C{r}", puntos, h="center")
            if paginas:
                # reparto sugerido de las páginas A4 según los puntos de cada apartado (en medias páginas);
                # se puede sobrescribir
                _escribe(ws, f"D{r}", f'=IF(OR($D${cab}="",C{r}="",SUM($C${primera}:$C${ultima})=0),"",'
                                      f'ROUND($D${cab}*C{r}/SUM($C${primera}:$C${ultima})*2,0)/2)', h="center")
            else:
                _escribe(ws, f"D{r}")
            _escribe(ws, f"E{r}", h="center")
            _escribe(ws, f"F{r}")
            _escribe(ws, f"G{r}", h="center")
            _casilla(ws, f"H{r}")
            _escribe(ws, f"I{r}", tam=10)
            _casilla(ws, f"J{r}")
            _escribe(ws, f"K{r}", tam=10)
            ws.row_dimensions[r].height = _alto((nombre, anch["B"]), minimo=22)
            r += 1
        suma = r
        _escribe(ws, f"A{suma}", relleno=GRIS)
        _escribe(ws, f"B{suma}", "Suma de apartados", cursiva=True, h="right", relleno=GRIS, tam=10)
        for col in "CDEG":
            _escribe(ws, f"{col}{suma}", f"=SUM({col}{primera}:{col}{ultima})", relleno=GRIS, h="center", negrita=True, tam=10)
        for col in "FHIJK":
            _escribe(ws, f"{col}{suma}", relleno=GRIS)
        # Avisos en rojo si los apartados no cuadran con el límite del pliego o si te pasas
        if ws[f"C{cab}"].value is not None:
            ws.conditional_formatting.add(f"C{suma}", FormulaRule(formula=[f"AND(C{suma}>0,C{suma}<>$C${cab})"], font=rojo))
            ws.conditional_formatting.add(f"G{suma}", FormulaRule(formula=[f"G{suma}>$C${cab}"], font=rojo))
        if paginas and ws[f"D{cab}"].value is not None:
            ws.conditional_formatting.add(f"D{suma}", FormulaRule(formula=[f"D{suma}>$D${cab}"], font=rojo))
        if paginas and ws[f"E{cab}"].value is not None:
            ws.conditional_formatting.add(f"E{suma}", FormulaRule(formula=[f"E{suma}>$E${cab}"], font=rojo))
        r += 1
        return suma

    # --- Juicio de valor
    fila_bloque_jv = titulo_bloque("CRITERIOS POR JUICIO DE VALOR")
    jv = _lista(datos.get("juicio_valor"))
    primera_cab = len(cabeceras_puntos)
    if not jv:
        _revisar(ws, f"A{r}:K{r}")
        r += 1
    for crit in jv:
        cab = cabecera_criterio(_txt(crit.get("nombre")), _num(crit.get("puntos")), _num(crit.get("paginas_a4")), _num(crit.get("paginas_a3")))
        partes = []
        if _txt(crit.get("formato")):
            partes.append(f"Formato: {_txt(crit.get('formato'))}")
        if _txt(crit.get("acreditacion")):
            partes.append(f"Acreditación: {_txt(crit.get('acreditacion'))}")
        nota("\n".join(partes))
        subs = [(_txt(a.get("nombre")), _num(a.get("puntos"))) for a in _lista(crit.get("apartados"))]
        apartados(cab, subs, paginas=True, extra_vacias=0 if subs else 3)
    r += 1

    # --- Fórmulas
    titulo_bloque("CRITERIOS POR FÓRMULA")
    formulas = _lista(datos.get("formulas"))
    if not formulas:
        _revisar(ws, f"A{r}:K{r}")
        r += 1
    for crit in formulas:
        cab = cabecera_criterio(_txt(crit.get("nombre")), _num(crit.get("puntos")))
        _escribe(ws, f"G{cab}", None, relleno=AZUL_SUAVE, h="center")
        form = _txt(crit.get("formula"))
        _escribe(ws, f"A{r}", "Fórmula", negrita=True, h="right")
        if form:
            _combina(ws, f"B{r}:K{r}", form, tam=11)
            ws.row_dimensions[r].height = _alto((form, ancho_total - anch["A"]), minimo=30)
        else:
            _revisar(ws, f"B{r}:K{r}")
        r += 1
        nota(f"Acreditación: {_txt(crit.get('acreditacion'))}" if _txt(crit.get("acreditacion")) else None)
    r += 1

    # --- Otros criterios automáticos
    titulo_bloque("OTROS CRITERIOS AUTOMÁTICOS (experiencia, mejoras, compromisos…)")
    otros = _lista(datos.get("otros"))
    if not otros:
        _combina(ws, f"A{r}:K{r}", "El pliego no recoge otros criterios automáticos (o no se han localizado).", cursiva=True, tam=10)
        r += 1
    for crit in otros:
        cab = cabecera_criterio(_txt(crit.get("nombre")), _num(crit.get("puntos")))
        nota(f"Acreditación: {_txt(crit.get('acreditacion'))}" if _txt(crit.get("acreditacion")) else None)
        subs = []
        for a in _lista(crit.get("apartados")):
            nombre, regla = _txt(a.get("nombre")) or "", _txt(a.get("regla"))
            subs.append((f"{nombre}\n{regla}" if regla else nombre, _num(a.get("puntos"))))
        apartados(cab, subs, paginas=False, extra_vacias=0 if subs else 1)
    r += 1

    umbral = _txt(datos.get("umbral")) or _txt((datos.get("info") or {}).get("umbral_tecnico") if isinstance(datos.get("info"), dict) else None)
    if umbral:
        _escribe(ws, f"A{r}", "Umbral mínimo", negrita=True, h="right", relleno=AZUL)
        _combina(ws, f"B{r}:K{r}", umbral)
        ws.row_dimensions[r].height = _alto((umbral, ancho_total - anch["A"]), minimo=22)
        r += 1

    # Total de puntos del pliego (suma de los límites de cada criterio)
    if cabeceras_puntos:
        _escribe(ws, "C1", "=" + "+".join(f"N(C{f})" for f in cabeceras_puntos), negrita=True, h="center",
                 num='0" ptos"', relleno=AZUL_SUAVE)
    ws["C1"].border = BORDE


# ------------------------------------------------------------------ Sobres
def _hoja_sobres(wb, bloque5, exp):
    ws = wb.create_sheet("Sobres")
    _preparar_hoja(ws)
    anch = {"A": 5, "B": 46, "C": 72, "D": 16, "E": 9, "F": 36}
    _anchos(ws, anch)
    _combina(ws, "A1:D1", f"Sobres a entregar · {exp}", negrita=True, tam=14, color=TEAL, borde=False)
    ws.row_dimensions[1].height = 30
    for col, texto in zip("ABCDEF", ["Nº", "Documento", "Qué debe contener (según el pliego)", "Responsable", "Hecho", "Notas"]):
        _escribe(ws, f"{col}2", texto, negrita=True, relleno=AZUL, h="center")
    ws.freeze_panes = "A3"
    r = 3
    rangos_entrega = []

    def bloque(titulo, documentos, notas, relleno, entrega=True):
        nonlocal r
        cab = r
        _combina(ws, f"A{r}:D{r}", titulo, negrita=True, relleno=relleno, tam=12)
        ws.row_dimensions[r].height = _alto((titulo, sum(list(anch.values())[:4])), minimo=22)
        primera, ultima = r + 1, r + len(documentos)
        if documentos:
            _escribe(ws, f"E{r}", f'=COUNTIF(E{primera}:E{ultima},TRUE)&"/"&ROWS(E{primera}:E{ultima})', negrita=True, relleno=relleno, h="center")
        else:
            _escribe(ws, f"E{r}", None, relleno=relleno)
        _escribe(ws, f"F{r}", None, relleno=relleno)
        r += 1
        for i, doc in enumerate(documentos, 1):
            detalle = "\n".join(f"• {d}" for d in doc.get("detalle", []) if d)
            _escribe(ws, f"A{r}", i, h="center")
            _escribe(ws, f"B{r}", doc["nombre"], negrita=True)
            _escribe(ws, f"C{r}", detalle, tam=10)
            _escribe(ws, f"D{r}")
            _casilla(ws, f"E{r}")
            _escribe(ws, f"F{r}", tam=10)
            ws.row_dimensions[r].height = _alto((doc["nombre"], anch["B"]), (detalle, anch["C"] * 1.1), minimo=22)
            r += 1
        if notas:
            texto = "Condiciones y avisos: " + " · ".join(notas)
            _combina(ws, f"A{r}:F{r}", texto, cursiva=True, tam=10, relleno=GRIS_CLARO)
            ws.row_dimensions[r].height = _alto((texto, sum(anch.values())), minimo=18)
            r += 1
        if entrega and documentos:
            rangos_entrega.append((primera, ultima))
        r += 1

    for sobre in bloque5["sobres"]:
        bloque(sobre["nombre"], sobre["documentos"], sobre["notas"], AZUL)
    if bloque5["posterior"]:
        bloque("Posterior adjudicatario · solo si somos el propuesto como adjudicatario (no se entrega ahora)",
               bloque5["posterior"], bloque5.get("notas_posterior"), GRIS, entrega=False)
    if rangos_entrega:
        hechos = "+".join(f"COUNTIF(E{a}:E{b},TRUE)" for a, b in rangos_entrega)
        total = "+".join(f"ROWS(E{a}:E{b})" for a, b in rangos_entrega)
        _combina(ws, "E1:F1", f'="Presentación: "&({hechos})&"/"&({total})', negrita=True, h="center", relleno=AZUL_SUAVE)


# ------------------------------------------------------------------ casillas nativas de Excel 365
_FEATURE_PROPERTY_BAG = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<FeaturePropertyBags xmlns="http://schemas.microsoft.com/office/spreadsheetml/2022/featurepropertybag">'
    '<bag type="Checkbox"/><bag type="XFControls"><bagId k="CellControl">0</bagId></bag>'
    '<bag type="XFComplement"><bagId k="XFControls">1</bagId></bag>'
    '<bag type="XFComplements" extRef="XFComplementsMapperExtRef"><a k="MappedFeaturePropertyBags"><bagId>2</bagId></a></bag>'
    '</FeaturePropertyBags>'
)
_EXT_CASILLA = ('<extLst><ext uri="{C7286773-470A-42A8-94C5-96B5CB345126}" '
                'xmlns:xfpb="http://schemas.microsoft.com/office/spreadsheetml/2022/featurepropertybag">'
                '<xfpb:xfComplement i="0"/></ext></extLst>')


def _activar_casillas(ruta):
    """openpyxl solo escribe VERDADERO/FALSO. Para que Excel 365 los muestre como casillas se añade lo mismo
    que guarda el propio Excel: el estilo de celda marcado como «casilla» (extLst/xfComplement) y la parte
    featurePropertyBag. Las celdas con estilo «shrinkToFit» son las casillas (ningún otro estilo lo usa)."""
    origen = Path(ruta)
    tmp = origen.with_suffix(".tmp")
    with zipfile.ZipFile(origen) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            datos = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                s = datos.decode("utf-8")
                ini = s.index("<cellXfs")
                fin = s.index("</cellXfs>")
                bloque = re.sub(r'(<xf\b(?:(?!</xf>).)*?shrinkToFit="1"(?:(?!</xf>).)*?)</xf>',
                                lambda m: m.group(1) + _EXT_CASILLA + "</xf>", s[ini:fin], flags=re.DOTALL)
                datos = (s[:ini] + bloque + s[fin:]).encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                s = datos.decode("utf-8")
                s = s.replace("</Types>", '<Override PartName="/xl/featurePropertyBag/featurePropertyBag.xml" '
                                          'ContentType="application/vnd.ms-excel.featurepropertybag+xml"/></Types>')
                datos = s.encode("utf-8")
            elif item.filename == "xl/_rels/workbook.xml.rels":
                s = datos.decode("utf-8")
                s = s.replace("</Relationships>", '<Relationship Id="rIdFPB1" Type="http://schemas.microsoft.com/office/2022/11/'
                                                  'relationships/FeaturePropertyBag" Target="featurePropertyBag/featurePropertyBag.xml"/>'
                                                  '</Relationships>')
                datos = s.encode("utf-8")
            zout.writestr(item, datos)
        zout.writestr("xl/featurePropertyBag/featurePropertyBag.xml", _FEATURE_PROPERTY_BAG)
    shutil.move(str(tmp), str(origen))


# ------------------------------------------------------------------ API
def crear_excel(registro, revision, bloque5, destino):
    """Crea el Excel de la licitación en `destino` (ruta de un .xlsx). Devuelve la ruta."""
    datos = revision.get("datos_excel") if isinstance(revision.get("datos_excel"), dict) else {}
    exp = str(registro.get("folder_id") or "")
    wb = Workbook()
    wb.properties.title = f"Ficha de la licitación {exp}"
    wb.properties.creator = "Monitor de licitaciones · CAI Consultores"
    _hoja_info(wb, registro, datos, exp, datetime.now())
    _hoja_solvencia(wb, datos, exp)
    _hoja_adjudicacion(wb, datos, exp)
    _hoja_sobres(wb, bloque5, exp)
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d) / "excel.xlsx"
        wb.save(tmp)
        _activar_casillas(tmp)
        shutil.copyfile(tmp, destino)
    return destino
